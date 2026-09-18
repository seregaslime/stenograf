"""Документы базы знаний: приём файла, текст из него, свои и чужие документы.

Пункт 5а. Человек загружает txt, md, docx или pdf — регламент, ТЗ, заметки, — и
документ ищется тем же поиском, что и встречи. Храним извлечённый текст, а не
файл: искать нужно по тексту, оригинал у человека и так есть.
"""
import re
import zipfile
from datetime import datetime
from io import BytesIO
from pathlib import PurePath
from typing import Optional

import docx
from pypdf import PdfReader
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from .db.models import Document

ALLOWED_SUFFIXES = (".txt", ".md", ".docx", ".pdf")
# Предел на ТЕКСТ — не про место в базе, а про время индексации: векторы считает
# модель на машине человека, bge-m3 на ПК — 4.1 куска в секунду (замер
# 17.09.2026). Мегабайт текста — около 1700 кусков по 600 символов, то есть ~7
# минут, пока поиск ждёт индексацию. Больше этого человек сочтёт зависанием.
MAX_BYTES = 1_000_000
СЛИШКОМ_МНОГО_ТЕКСТА = ("В файле больше миллиона символов: "
                        "его индексация заняла бы больше семи минут.")
# Предел на ФАЙЛ. У txt и md файл и есть текст, а docx и pdf носят с собой
# картинки, шрифты и разметку: на мегабайт текста файл легко весит десять.
MAX_FILE_BYTES = {".docx": 20_000_000, ".pdf": 20_000_000}
# Сколько может весить содержимое архива docx в распакованном виде. Архив на
# мегабайт из одних нулей распаковывается в гигабайты и кладёт сервер на лопатки
# ещё до того, как мы дойдём до текста, — это называется «zip-бомба».
MAX_UNPACKED_BYTES = 100_000_000
# Сколько страниц pdf разбираем. Предел на символы останавливает обычный
# текстовый документ сам, но у скана страницы пустые: миллион символов не
# наберётся никогда, а время на разбор каждой всё равно уходит. Тысяча страниц —
# заведомо больше любого регламента и заведомо меньше того, на чём сервер
# задумается надолго.
MAX_PDF_PAGES = 1000


class DocumentRejected(ValueError):
    """Файл не принят. status — код ответа, message — что показать человеку."""

    def __init__(self, status: int, message: str):
        super().__init__(message)
        self.status = status
        self.message = message


def decode_text(raw: bytes) -> str:
    """Текст файла: UTF-16 с меткой, UTF-8, а если не читается — cp1251.

    cp1251 — не догадка на всякий случай: Блокнот и Word в русской Windows
    сохраняют txt именно в ней, а у куратора Windows. UTF-8 проверяется раньше,
    потому что байты русского UTF-8 в cp1251 тоже «прочитаются» — кракозябрами.
    UTF-16 — это «Юникод» в Блокноте; узнаётся по метке в начале файла. Первая
    версия его отбивала по нулевым байтам и промахнулась: у русских букв в
    UTF-16 нулевых байтов нет, и текст без пробелов ушёл бы в cp1251 мусором.

    Нулевой байт вне UTF-16 — признак не текста: pdf с переименованным
    расширением. Такое честнее отбить с подсказкой, чем сохранить мусор, который
    потом найдётся в поиске.
    """
    if raw.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return raw.decode("utf-16").replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError:
            raise DocumentRejected(415, "Не удалось прочитать текст: сохраните файл в кодировке UTF-8.")
    if b"\x00" in raw:
        raise DocumentRejected(415, "Это не текстовый файл.")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = raw.decode("cp1251")
        except UnicodeDecodeError:
            raise DocumentRejected(415, "Не удалось прочитать текст: сохраните файл в кодировке UTF-8.")
    return text.replace("\r\n", "\n").replace("\r", "\n")


def docx_text(raw: bytes) -> str:
    """Текст из docx: абзацы и таблицы, по порядку.

    Таблицы вытаскиваются отдельно: в документе они лежат не среди абзацев, и
    без этого из регламента с таблицей сроков в поиск попала бы одна вода
    вокруг неё. Ячейки строки склеиваются через « | » — так строка таблицы
    остаётся одной строкой текста, а не рассыпается на слова.

    Сноски, колонтитулы и надписи в фигурах не берём: в них обычно номера
    страниц и реквизиты, для поиска это шум.
    """
    проверить_архив(raw)
    try:
        документ = docx.Document(BytesIO(raw))
    except Exception:
        raise DocumentRejected(415, "Не удалось прочитать docx: файл повреждён или это не Word.")

    куски: list[str] = []
    абзацы = {абзац._element: абзац for абзац in документ.paragraphs}
    таблицы = {таблица._element: таблица for таблица in документ.tables}
    for элемент in документ.element.body.iterchildren():
        if элемент in абзацы:
            текст = абзацы[элемент].text.strip()
            if текст:
                куски.append(текст)
        elif элемент in таблицы:
            for строка in таблицы[элемент].rows:
                ячейки = [я.text.strip() for я in строка.cells if я.text.strip()]
                if ячейки:
                    куски.append(" | ".join(ячейки))
    return "\n\n".join(куски)


def pdf_text(raw: bytes) -> str:
    """Текст из pdf — постранично, со склейкой слов, разорванных переносом.

    В pdf нет абзацев: есть буквы с координатами, и строки собирает уже pypdf.
    Поэтому склеиваем переносы («доку-\\nмент» → «документ»), иначе половинки
    слов попадут в вектор как отдельные слова, а страницы разделяем пустой
    строкой — по ней document_chunks режет текст на куски.

    Колонтитулы и номера страниц не вычищаем: отличить их от текста без замера
    нельзя, а лишняя строка «стр. 4» в куске поиску почти не мешает.

    Сканы отбиваем с объяснением. В таком pdf текста нет вовсе — только
    картинка страницы, — и сохранить его молча значило бы завести документ,
    который никогда ничего не найдёт.
    """
    try:
        reader = PdfReader(BytesIO(raw))
        зашифрован = reader.is_encrypted
    except Exception:
        raise DocumentRejected(415, "Не удалось прочитать pdf: файл повреждён.")
    if зашифрован:
        # Пустой пароль pypdf подбирает сам — им закрыты pdf «только от правки».
        # Всё остальное (не подошёл пароль, незнакомый способ шифрования) — для
        # человека одно и то же: файл нужно открыть самому и пересохранить.
        try:
            открылся = reader.decrypt("")
        except Exception:
            открылся = False
        if not открылся:
            raise DocumentRejected(415, "Pdf защищён паролем: снимите пароль и загрузите заново.")
    try:
        страницы = reader.pages
    except Exception:
        raise DocumentRejected(415, "Не удалось прочитать pdf: файл повреждён.")
    if len(страницы) > MAX_PDF_PAGES:
        raise DocumentRejected(413, f"В pdf больше {MAX_PDF_PAGES} страниц.")

    куски: list[str] = []
    символов = 0
    for страница in страницы:
        try:
            текст = страница.extract_text() or ""
        except Exception:
            raise DocumentRejected(415, "Не удалось прочитать pdf: файл повреждён.")
        текст = re.sub(r"(\w)-\n(\w)", r"\1\2", текст)
        # Точки-заполнители из оглавления («Введение .......... 5»): в обычном
        # тексте четыре точки подряд не встречаются, а строка оглавления состоит
        # из них наполовину и вытесняет из куска настоящие слова
        текст = re.sub(r"[ \t]*\.{4,}[ \t]*", " ", текст).strip()
        if not текст:
            continue
        куски.append(текст)
        # Считаем по ходу, а не в конце: миллион символов набирается страниц за
        # триста, и разбирать ради отказа остальные семьсот незачем
        символов += len(текст)
        if символов > MAX_BYTES:
            raise DocumentRejected(413, СЛИШКОМ_МНОГО_ТЕКСТА)
    if not куски:
        raise DocumentRejected(415, "В pdf нет текстового слоя: похоже, это скан. "
                                    "Распознайте его или загрузите текстовую версию.")
    return "\n\n".join(куски)


def проверить_архив(raw: bytes) -> None:
    """Архив не должен распаковываться в гигабайты (zip-бомба)."""
    try:
        with zipfile.ZipFile(BytesIO(raw)) as архив:
            распакованный = sum(файл.file_size for файл in архив.infolist())
    except zipfile.BadZipFile:
        raise DocumentRejected(415, "Не удалось прочитать docx: файл повреждён или это не Word.")
    if распакованный > MAX_UNPACKED_BYTES:
        raise DocumentRejected(413, "Файл распаковывается в слишком большой документ.")


def title_from_filename(filename: str) -> str:
    """Название — имя файла без расширения: так человек узнает его в выдаче."""
    return PurePath(filename).stem.strip()[:300] or "Документ"


def check_filename(filename: str) -> str:
    """Расширение файла, если оно нам подходит. Иначе — отказ с понятным текстом."""
    суффикс = PurePath(filename).suffix.lower()
    if суффикс not in ALLOWED_SUFFIXES:
        raise DocumentRejected(415, "Пока принимаются файлы .txt, .md, .docx и .pdf.")
    return суффикс


def извлечь_текст(суффикс: str, raw: bytes) -> str:
    return {".docx": docx_text, ".pdf": pdf_text}.get(суффикс, decode_text)(raw)


def create(db: Session, filename: str, raw: bytes, owner_id: Optional[int]) -> Document:
    суффикс = check_filename(filename)
    предел = MAX_FILE_BYTES.get(суффикс, MAX_BYTES)
    if len(raw) > предел:
        raise DocumentRejected(413, f"Файл больше {предел // 1_000_000} МБ.")
    text = извлечь_текст(суффикс, raw).strip()
    if not text:
        raise DocumentRejected(400, "В файле нет текста.")
    if len(text) > MAX_BYTES:
        raise DocumentRejected(413, СЛИШКОМ_МНОГО_ТЕКСТА)
    document = Document(owner_id=owner_id, title=title_from_filename(filename), text=text)
    db.add(document)
    db.flush()
    return document


def for_owner(db: Session, document_id: int, owner_id: Optional[int]) -> Optional[Document]:
    """Документ, если он есть и принадлежит этому человеку. Чужой неотличим от
    несуществующего — как у встреч и голосов."""
    document = db.get(Document, document_id)
    if document is None or (owner_id is not None and document.owner_id != owner_id):
        return None
    return document


def list_for_owner(db: Session, owner_id: Optional[int]) -> list[dict]:
    query = select(Document.id, Document.title, Document.created_at, func.length(Document.text))
    if owner_id is not None:
        query = query.where(Document.owner_id == owner_id)
    return [to_dict(id_, title, created_at, chars)
            for id_, title, created_at, chars in db.execute(query.order_by(Document.created_at.desc()))]


def to_dict(id_: int, title: str, created_at: datetime, chars: int) -> dict:
    return {"id": id_, "title": title, "created_at": created_at.isoformat(), "chars": chars}
